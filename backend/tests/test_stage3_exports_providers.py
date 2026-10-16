from io import BytesIO
import unittest
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pymupdf as fitz
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from app.config.settings import settings
from app.rag.llm import ProviderConfigurationError, get_llm, provider_status
from app.rag.embeddings import get_embedding_model
from app.schemas.export import ExportRequest
from app.services.export_service import create_excel_export, create_pdf_export
from app.rag import vector_store
from chromadb import PersistentClient
from app.main import app


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.request = ExportRequest.model_validate({
            "title": "Stage 3 verification",
            "messages": [
                {"role": "user", "content": "What is the retention period?", "created_at": "2026-10-11T08:30:00Z"},
                {"role": "assistant", "content": "Records are kept for 7 years.", "sources": [{"source": "Policy.pdf", "page_number": 4, "chunk_index": 2, "url": "https://example.org/policy"}]},
            ],
        })

    def test_pdf_export_is_readable_and_contains_answer_and_citation(self):
        payload = create_pdf_export(self.request)
        pdf = fitz.open(stream=payload, filetype="pdf")
        text = "\n".join(page.get_text() for page in pdf)
        self.assertIn("Stage 3 verification", text)
        self.assertIn("Records are kept for 7 years.", text)
        self.assertIn("Policy.pdf", text)
        self.assertIn("Page 4", text)

    def test_excel_export_is_readable_and_contains_answer_and_citation(self):
        workbook = load_workbook(BytesIO(create_excel_export(self.request)), read_only=True)
        rows = list(workbook.active.values)
        self.assertEqual(rows[0], ("Role", "Content", "Timestamp", "Sources"))
        self.assertEqual(rows[2][1], "Records are kept for 7 years.")
        self.assertIn("Page 4", rows[2][3])
        self.assertIn("https://example.org/policy", rows[2][3])

    def test_export_api_returns_downloadable_pdf_and_excel(self):
        client = TestClient(app)
        payload = self.request.model_dump(mode="json")
        pdf = client.post("/api/v1/exports/pdf", json=payload)
        excel = client.post("/api/v1/exports/excel", json=payload)
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.headers["content-type"].startswith("application/pdf"))
        self.assertTrue(pdf.headers["content-disposition"].endswith('"pharma-conversation.pdf"'))
        self.assertEqual(excel.status_code, 200)
        self.assertIn("spreadsheetml.sheet", excel.headers["content-type"])
        self.assertEqual(load_workbook(BytesIO(excel.content), read_only=True).active.max_row, 3)


class ProviderTests(unittest.TestCase):
    def test_ollama_is_default_and_status_does_not_expose_secrets(self):
        with patch.object(settings, "llm_provider", "ollama"), patch.object(settings, "embedding_provider", "ollama"), patch.object(settings, "openai_api_key", "super-secret"):
            status = provider_status()
        self.assertEqual(status["provider"], "ollama")
        self.assertEqual(status["embedding_model"], "nomic-embed-text:latest")
        self.assertNotIn("api_key", status)
        self.assertNotIn("super-secret", str(status))
        self.assertFalse(status["reindex_required"])

    def test_openai_requires_server_side_key_without_calling_paid_service(self):
        with patch.object(settings, "llm_provider", "openai"), patch.object(settings, "openai_api_key", None):
            with self.assertRaises(ProviderConfigurationError):
                get_llm()

    def test_openai_factories_select_configured_clients_without_network_calls(self):
        from langchain_openai import ChatOpenAI, OpenAIEmbeddings
        with patch.object(settings, "llm_provider", "openai"), patch.object(settings, "embedding_provider", "openai"), patch.object(settings, "openai_api_key", "test-key"):
            self.assertIsInstance(get_llm(), ChatOpenAI)
            self.assertIsInstance(get_embedding_model(), OpenAIEmbeddings)

    def test_new_embedding_provider_reports_separate_reindex_collection(self):
        with patch.object(settings, "llm_provider", "ollama"), patch.object(settings, "embedding_provider", "openai"), patch.object(settings, "openai_api_key", "test-key"):
            status = provider_status()
        self.assertTrue(status["reindex_required"])
        self.assertNotEqual(status["vector_collection"], "pharma_documents")


class VectorCleanupTests(unittest.TestCase):
    def test_document_cleanup_covers_saved_embedding_spaces_only(self):
        with TemporaryDirectory() as directory:
            client = PersistentClient(path=directory)
            try:
                for name in ("pharma_documents", "pharma_documents_openai_test", "unrelated_collection"):
                    collection = client.create_collection(name)
                    collection.add(
                        ids=[f"{name}-target", f"{name}-other"],
                        embeddings=[[0.1, 0.2], [0.2, 0.3]],
                        documents=["target chunk", "other document chunk"],
                        metadatas=[
                            {"document_id": "target", "content_hash": "target-hash"},
                            {"document_id": "other", "content_hash": "other-hash"},
                        ],
                    )
                with patch.object(vector_store, "CHROMA_PATH", directory), patch.object(
                    vector_store, "PersistentClient", return_value=client
                ):
                    vector_store.delete_document_vectors(document_id="target")
                self.assertEqual(client.get_collection("pharma_documents").count(), 1)
                self.assertEqual(client.get_collection("pharma_documents_openai_test").count(), 1)
                self.assertEqual(client.get_collection("unrelated_collection").count(), 2)
            finally:
                client._system.stop()


if __name__ == "__main__":
    unittest.main()
