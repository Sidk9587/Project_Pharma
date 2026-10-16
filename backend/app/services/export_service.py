"""In-memory PDF and spreadsheet conversation exports."""

from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from app.schemas.export import ExportRequest


def _source_text(source: dict) -> str:
    parts = [str(source.get(key)) for key in ("source", "url", "domain") if source.get(key)]
    if source.get("page_number") is not None:
        parts.append(f"Page {source['page_number']}")
    return " | ".join(parts)


def _spreadsheet_text(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def create_pdf_export(request: ExportRequest) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.65*inch, leftMargin=0.65*inch)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ExportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#145b56"))
    body = ParagraphStyle("MessageBody", parent=styles["BodyText"], leading=15, spaceAfter=8)
    story = [Paragraph(escape(request.title), title), Spacer(1, 16)]
    for message in request.messages:
        stamp = f" · {message.created_at.isoformat()}" if message.created_at else ""
        story.append(Paragraph(f"<b>{escape(message.role.title()+stamp)}</b>", styles["Heading3"]))
        for line in message.content.splitlines() or [message.content]:
            story.append(Paragraph(escape(line) or "&nbsp;", body))
        for source in message.sources:
            label = _source_text(source)
            if label:
                story.append(Paragraph(f"<font color='#51636c'>Source: {escape(label)}</font>", body))
        story.append(Spacer(1, 8))
    doc.build(story)
    return output.getvalue()


def create_excel_export(request: ExportRequest) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conversation"
    sheet.append(["Role", "Content", "Timestamp", "Sources"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="145B56")
    for message in request.messages:
        sheet.append([message.role, _spreadsheet_text(message.content), message.created_at.isoformat() if message.created_at else "", _spreadsheet_text("\n".join(filter(None, (_source_text(source) for source in message.sources))))])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 14, "B": 80, "C": 28, "D": 70}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
